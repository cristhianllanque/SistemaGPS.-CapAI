import io
from fastapi import APIRouter, Depends, HTTPException, Response, Query
from sqlalchemy.orm import Session
from datetime import datetime
import openpyxl
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet

from app.api.deps import get_db, get_current_user
from app.db import models
from app.api.vehiculos import haversine

router = APIRouter()

@router.get("/consolidado")
def get_reporte_consolidado(
    fecha_inicio: str = Query(...),
    fecha_fin: str = Query(...),
    formato: str = Query(..., description="excel, pdf, o json"),
    db: Session = Depends(get_db),
    current_user: models.Usuario = Depends(get_current_user)
):
    if formato not in ['excel', 'pdf', 'json']:
        raise HTTPException(status_code=400, detail="Formato inválido. Use 'excel', 'pdf', o 'json'.")

    try:
        inicio = datetime.fromisoformat(fecha_inicio.replace('Z', '+00:00'))
        fin = datetime.fromisoformat(fecha_fin.replace('Z', '+00:00'))
    except ValueError:
        raise HTTPException(status_code=400, detail="Formato de fecha inválido")

    # Get vehicles
    vehiculos = db.query(models.Vehiculo).filter(
        models.Vehiculo.empresa_id == current_user.empresa_id,
        models.Vehiculo.activo == True
    ).all()

    report_data = []

    for v in vehiculos:
        # Get GPS readings
        lecturas = db.query(models.LecturaGPS).filter(
            models.LecturaGPS.vehiculo_id == v.id,
            models.LecturaGPS.timestamp >= inicio,
            models.LecturaGPS.timestamp <= fin
        ).order_by(models.LecturaGPS.timestamp).all()

        distancia_total = 0.0
        velocidad_maxima = 0
        tiempo_encendido_min = 0.0

        for i, lec in enumerate(lecturas):
            if lec.velocidad > velocidad_maxima:
                velocidad_maxima = lec.velocidad
                
            if i > 0:
                prev_lec = lecturas[i-1]
                dist = haversine(prev_lec.latitud, prev_lec.longitud, lec.latitud, lec.longitud)
                distancia_total += dist
                
                dt_mins = (lec.timestamp - prev_lec.timestamp).total_seconds() / 60.0
                if prev_lec.ignicion:
                    tiempo_encendido_min += dt_mins

        # Get Panic alerts
        alertas_count = db.query(models.AlertaPanico).filter(
            models.AlertaPanico.vehiculo_id == v.id,
            models.AlertaPanico.fecha_hora >= inicio,
            models.AlertaPanico.fecha_hora <= fin
        ).count()

        conductor_nombre = v.conductor.nombre if v.conductor else "Sin asignar"

        report_data.append({
            "Placa": v.placa,
            "Conductor": conductor_nombre,
            "Km Totales": round(distancia_total, 2),
            "Horas Motor": round(tiempo_encendido_min / 60.0, 2),
            "Velocidad Max (km/h)": velocidad_maxima,
            "Alertas Pánico": alertas_count
        })

    if formato == 'excel':
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Reporte Flota"

        headers = ["Placa", "Conductor", "Km Totales", "Horas Motor", "Velocidad Max (km/h)", "Alertas Pánico"]
        ws.append(headers)

        for row in report_data:
            ws.append([
                row["Placa"],
                row["Conductor"],
                row["Km Totales"],
                row["Horas Motor"],
                row["Velocidad Max (km/h)"],
                row["Alertas Pánico"]
            ])

        # Formatting headers
        for cell in ws[1]:
            cell.font = openpyxl.styles.Font(bold=True)
            cell.fill = openpyxl.styles.PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
            cell.font = openpyxl.styles.Font(color="FFFFFF", bold=True)

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)

        headers = {
            'Content-Disposition': 'attachment; filename="reporte_flota.xlsx"'
        }
        return Response(content=buffer.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers=headers)

    elif formato == 'pdf':
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=letter)
        styles = getSampleStyleSheet()
        elements = []

        elements.append(Paragraph("Reporte Consolidado de Flota", styles['Title']))
        elements.append(Paragraph(f"Empresa ID: {current_user.empresa_id}", styles['Normal']))
        elements.append(Paragraph(f"Periodo: {inicio.strftime('%d/%m/%Y %H:%M')} a {fin.strftime('%d/%m/%Y %H:%M')}", styles['Normal']))
        elements.append(Spacer(1, 20))

        table_data = [["Placa", "Conductor", "Km Totales", "Horas Motor", "Velocidad Max", "Alertas Pánico"]]
        for row in report_data:
            table_data.append([
                row["Placa"],
                row["Conductor"],
                str(row["Km Totales"]),
                str(row["Horas Motor"]),
                str(row["Velocidad Max (km/h)"]),
                str(row["Alertas Pánico"])
            ])

        t = Table(table_data)
        t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#4F46E5')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
            ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
            ('GRID', (0, 0), (-1, -1), 1, colors.black),
        ]))
        elements.append(t)

        doc.build(elements)
        buffer.seek(0)

        headers = {
            'Content-Disposition': 'attachment; filename="reporte_flota.pdf"'
        }
        return Response(content=buffer.getvalue(), media_type="application/pdf", headers=headers)

    elif formato == 'json':
        return {"data": report_data}
